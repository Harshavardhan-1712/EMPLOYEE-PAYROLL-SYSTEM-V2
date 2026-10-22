import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AuditEvent
from app.services.exports import safe_text
from tests.test_payroll_runs_api import emp, make_run
from tests.world import build, finalize, money, tot


@pytest.fixture()
def world(login):
    ids, runs, exp = build(login)
    return ids, runs, exp, login("hr", "hr_editor")


def parse_csv(resp):
    assert resp.content.startswith(b"\xef\xbb\xbf"), "UTF-8 BOM expected so Excel reads names correctly"
    assert resp.headers["content-type"] == "text/csv; charset=utf-8"
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))


def wb_of(resp):
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return load_workbook(io.BytesIO(resp.content))


@pytest.mark.parametrize("raw,expected", [("=1+1", "'=1+1"), ("+91 98", "'+91 98"), ("-2+3", "'-2+3"), ("@SUM(A1)", "'@SUM(A1)"), ("\t=1", "'\t=1"),
                                          ("\r=1", "'\r=1"), ("  =cmd", "'  =cmd"), ("Priya", "Priya"), ("A-1", "A-1"), ("", ""), (5, 5)])
def test_safe_text(raw, expected):
    assert safe_text(raw) == expected


# ------------------------------------------------------------------------------- payroll register CSV
def test_payroll_csv_rows_and_total_row_match_snapshot(world):
    ids, runs, exp, hr = world
    r = hr.get(f"/api/reports/payroll.csv?run_id={runs['2026-06']}")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == f'attachment; filename="payroll_register_2026-06_run{runs["2026-06"]}.csv"'
    assert r.headers["cache-control"] == "private, no-store" and r.headers["x-content-type-options"] == "nosniff"
    rows = parse_csv(r)
    head, body, total = rows[0], rows[1:-1], rows[-1]
    assert head[:4] == ["Employee ID", "Name", "Department", "Designation"] and len(body) == 3
    col = {h: i for i, h in enumerate(head)}
    for line, e in zip(body, exp["2026-06"]):
        assert line[col["Employee ID"]] == e["code"] and line[col["Gross salary"]] == money(e["gross"]) and line[col["Net salary"]] == money(e["net"])
        assert line[col["Overtime pay"]] == money(e["ot_pay"]) and line[col["Bonus"]] == money(e["bonus"])
    # total row recomputed from the exported rows themselves, and compared with the run's stored totals
    for h in ("Basic salary", "Bonus", "Overtime pay", "Gross salary", "Deductions", "Net salary"):
        assert Decimal(total[col[h]]) == sum(Decimal(b[col[h]]) for b in body), h
    run = hr.get(f"/api/payroll-runs/{runs['2026-06']}").json()
    assert total[0] == "TOTAL" and total[col["Net salary"]] == run["total_net"] and total[col["Gross salary"]] == run["total_gross"]


def test_csv_neutralises_formula_injection_but_leaves_numbers_alone(login):
    hr = login("hr", "hr_editor")
    evil = ['=HYPERLINK("http://evil.test","click")', "+1+1", "@SUM(A1:A9)", "-2+3"]
    for n, name in enumerate(evil):
        emp(hr, f"X{n}", name=name, department="=cmd|' /C calc'!A0", designation="\t=1", basic_salary="1000", bonus="0", deductions="3000", overtime_hours="0")
    rid = make_run(hr, "2026-05")["id"]
    finalize(login, rid)
    hr = login("hr", "hr_editor")
    rows = parse_csv(hr.get(f"/api/reports/payroll.csv?run_id={rid}"))
    names = [r[1] for r in rows[1:-1]]
    assert sorted(names) == sorted("'" + n for n in evil)
    assert all(r[2].startswith("'=") and r[3].startswith("'") for r in rows[1:-1])
    assert all(r[11] == "-2000.00" for r in rows[1:-1])                         # negative net is a number: not prefixed
    for path in (f"payroll.csv?run_id={rid}", f"department.csv?run_id={rid}", "employees.csv", "analytics.csv"):
        for line in parse_csv(hr.get(f"/api/reports/{path}")):
            for cell in line:
                assert not cell.startswith(("=", "+", "@", "\t", "\r")), (path, cell)    # only numeric '-' may start a cell


# ------------------------------------------------------------------------------- Excel workbook
def test_xlsx_structure_formatting_and_values(world):
    ids, runs, exp, hr = world
    rid = runs["2026-05"]
    r = hr.get(f"/api/reports/payroll.xlsx?run_id={rid}")
    assert r.status_code == 200 and r.headers["content-disposition"] == f'attachment; filename="payroll_report_2026-05_run{rid}.xlsx"'
    wb = wb_of(r)
    assert wb.sheetnames == ["Summary", "Payroll Register", "By Department"]
    run = hr.get(f"/api/payroll-runs/{rid}").json()

    s = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=4, max_col=2) if row[0].value}
    assert s["Period"] == "2026-05" and s["Employees"] == 3 and s["Status"] == "Finalized"
    assert Decimal(str(s["Total gross salary"])) == Decimal(run["total_gross"]) and Decimal(str(s["Total net salary"])) == Decimal(run["total_net"])
    assert isinstance(s["Total net salary"], (int, float, Decimal)), "totals must be numbers, not text"
    assert Decimal(str(s["Average net salary"])) == (tot(exp["2026-05"], "net") / 3).quantize(Decimal("0.01"))
    assert len(s["Exported (UTC)"]) == 19 and s["Exported (UTC)"][4] == "-"            # export timestamp present
    assert Decimal(str(s["Overtime rate (INR/hour)"])) == 200

    reg = wb["Payroll Register"]
    assert reg.freeze_panes == "C2"
    hdr = reg["A1"]
    assert hdr.font.bold and hdr.fill.fgColor.rgb.endswith("1F3A5F") and hdr.font.color.rgb.endswith("FFFFFF")
    assert reg.max_row == 1 + 3 + 1 and reg.cell(row=5, column=1).value == "TOTAL"
    assert all(reg.column_dimensions[c].width and reg.column_dimensions[c].width >= 10 for c in "ABCDEFGHIJKLM")
    assert reg.cell(row=2, column=12).number_format == '#,##0.00;[Red]-#,##0.00'
    assert reg.cell(row=3, column=1).fill.fgColor.rgb.endswith("FAF9F7") and reg.cell(row=2, column=1).fill.fgColor.rgb != reg.cell(row=3, column=1).fill.fgColor.rgb  # alternating rows
    assert Decimal(str(reg.cell(row=2, column=12).value)) == exp["2026-05"][0]["net"]
    assert Decimal(str(reg.cell(row=5, column=12).value)) == tot(exp["2026-05"], "net")

    dep = wb["By Department"]
    drows = {dep.cell(row=r, column=1).value: dep.cell(row=r, column=5).value for r in range(2, dep.max_row)}
    assert Decimal(str(drows["Eng"])) == tot([e for e in exp["2026-05"] if e["dept"] == "Eng"], "net")
    assert dep.cell(row=dep.max_row, column=1).value == "TOTAL"


def test_xlsx_never_contains_formula_cells(login):
    hr = login("hr", "hr_editor")
    emp(hr, "F1", name="=SUM(1+1)", department="@x", designation="+y", basic_salary="1000", bonus="0", deductions="0", overtime_hours="0")
    rid = make_run(hr, "2026-05")["id"]; finalize(login, rid)
    wb = wb_of(login("hr", "hr_editor").get(f"/api/reports/payroll.xlsx?run_id={rid}"))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", f"formula cell {ws.title}!{c.coordinate}: {c.value!r}"
    assert wb["Payroll Register"]["B2"].value == "'=SUM(1+1)" and wb["By Department"]["A2"].value == "'@x"


# ------------------------------------------------------------------------------- other reports
def test_department_csv_totals_match_run(world):
    ids, runs, exp, hr = world
    rid = runs["2026-05"]
    rows = parse_csv(hr.get(f"/api/reports/department.csv?run_id={rid}"))
    assert rows[0] == ["Department", "Employees", "Gross salary", "Deductions", "Net salary"]
    by = {r[0]: r for r in rows[1:-1]}
    assert by["Eng"][1] == "2" and by["Eng"][4] == money(tot([e for e in exp["2026-05"] if e["dept"] == "Eng"], "net")) and by["Ops"][1] == "1"
    run = hr.get(f"/api/payroll-runs/{rid}").json()
    assert rows[-1] == ["TOTAL", "3", run["total_gross"], run["total_deductions"], run["total_net"]]


def test_employee_register_filters(world):
    ids, runs, exp, hr = world
    allr = parse_csv(hr.get("/api/reports/employees.csv"))
    assert len(allr) == 1 + 3 and allr[0][0] == "Employee ID"
    inactive = parse_csv(hr.get("/api/reports/employees.csv?status=inactive"))
    assert [r[0] for r in inactive[1:]] == ["B1"]
    ops = parse_csv(hr.get("/api/reports/employees.csv?department=Ops"))
    assert [r[0] for r in ops[1:]] == ["B1"] and ops[1][6] == "30000.00"
    assert hr.get("/api/reports/employees.csv?status=bogus").status_code == 422
    assert 'filename="employee_register_inactive.csv"' in hr.get("/api/reports/employees.csv?status=inactive").headers["content-disposition"]


def test_analytics_csv_honours_filters(world):
    ids, runs, exp, hr = world
    rows = parse_csv(hr.get("/api/reports/analytics.csv"))
    assert rows[0] == ["Period", "Department", "Employees", "Gross salary", "Deductions", "Net salary", "Bonus", "Overtime pay"]
    assert [(r[0], r[1]) for r in rows[1:]] == [("2026-05", "Eng"), ("2026-05", "Ops"), ("2026-06", "Eng"), ("2026-06", "Ops"), ("2026-08", "Eng")]
    f = parse_csv(hr.get("/api/reports/analytics.csv?from=2026-06&to=2026-08&department=Eng"))
    assert [(r[0], r[1]) for r in f[1:]] == [("2026-06", "Eng"), ("2026-08", "Eng")]
    assert f[1][5] == money(tot([e for e in exp["2026-06"] if e["dept"] == "Eng"], "net"))
    assert hr.get("/api/reports/analytics.csv?from=2026-08&to=2026-05").status_code == 422
    assert 'filename="analytics_2026-06_to_2026-08.csv"' in hr.get("/api/reports/analytics.csv?from=2026-06&to=2026-08").headers["content-disposition"]


# ------------------------------------------------------------------------------- access, state, audit
def test_only_finalized_runs_can_be_exported(world, login):
    ids, runs, exp, hr = world
    draft = make_run(hr, "2026-09")["id"]
    for kind in ("payroll.csv", "payroll.xlsx", "department.csv"):
        r = hr.get(f"/api/reports/{kind}?run_id={draft}")
        assert r.status_code == 409 and "finalized" in r.json()["error"]["message"], kind
    login("ap", "payroll_approver").post(f"/api/payroll-runs/{runs['2026-05']}/reverse", json={"reason": "Deductions entered incorrectly"})
    hr = login("hr", "hr_editor")
    r = hr.get(f"/api/reports/payroll.xlsx?run_id={runs['2026-05']}")
    assert r.status_code == 409 and "reversed" in r.json()["error"]["message"]
    assert hr.get("/api/reports/payroll.csv?run_id=9999").status_code == 404
    assert hr.get("/api/reports/payroll.csv").status_code == 422
    assert hr.get("/api/reports/payroll.csv?run_id=abc").status_code == 422


def test_viewers_cannot_export_and_anonymous_is_rejected(world, login, client):
    ids, runs, exp, hr = world
    rid = runs["2026-05"]
    urls = [f"/api/reports/payroll.csv?run_id={rid}", f"/api/reports/payroll.xlsx?run_id={rid}", f"/api/reports/department.csv?run_id={rid}", "/api/reports/employees.csv", "/api/reports/analytics.csv"]
    v = login("vic", "viewer")
    assert [v.get(u).status_code for u in urls] == [403] * 5
    assert all(login("ap", "payroll_approver").get(u).status_code == 200 for u in urls)
    client.post("/api/auth/logout")
    assert [client.get(u).status_code for u in urls] == [401] * 5


def test_exports_are_audited_without_pay_figures(world, db):
    ids, runs, exp, hr = world
    rid = runs["2026-05"]
    hr.get(f"/api/reports/payroll.csv?run_id={rid}"); hr.get(f"/api/reports/payroll.xlsx?run_id={rid}"); hr.get("/api/reports/employees.csv?status=active")
    ev = db.scalars(select(AuditEvent).where(AuditEvent.action == "report.exported").order_by(AuditEvent.id)).all()
    assert [e.entity_id for e in ev] == ["payroll.csv", "payroll.xlsx", "employees.csv"]
    assert '"rows": 3' in ev[0].detail and f'"run_id": {rid}' in ev[0].detail and ev[0].actor_username == "hr"
    assert all("Amy" not in e.detail and "50000" not in e.detail for e in ev)
