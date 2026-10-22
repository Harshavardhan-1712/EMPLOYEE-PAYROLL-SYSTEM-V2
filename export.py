# ============================================================
# export.py — Export Payroll Reports
# Handles: CSV export, Excel export (openpyxl)
# ============================================================

import csv
import os
from datetime import datetime

import database as db
import payroll as pr

EXPORTS_DIR = os.path.join(os.path.dirname(__file__), "exports")
os.makedirs(EXPORTS_DIR, exist_ok=True)

FIELDNAMES = [
    "emp_id", "name", "department", "designation",
    "basic_salary", "bonus", "overtime_pay", "gross_salary",
    "deductions", "net_salary", "attendance", "overtime_hrs"
]


def _get_payroll_rows():
    """Return list of computed payroll dicts for all employees."""
    employees = db.get_all_employees()
    if not employees:
        print("\n⚠️  No data to export.")
        return []
    return [pr.compute_payroll_for_employee(emp) for emp in employees]


# ──────────────────────────────────────────────
# CSV Export
# ──────────────────────────────────────────────

def export_to_csv():
    """Export all employee payroll data to a CSV file."""
    rows = _get_payroll_rows()
    if not rows:
        return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename  = f"payroll_report_{timestamp}.csv"
    filepath  = os.path.join(EXPORTS_DIR, filename)

    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)

    print(f"\n✅  CSV exported  → {filepath}")


# ──────────────────────────────────────────────
# Excel Export (openpyxl)
# ──────────────────────────────────────────────

def export_to_excel():
    """Export all employee payroll data to an Excel (.xlsx) file."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("\n❌  openpyxl not installed. Run: pip install openpyxl")
        return

    rows = _get_payroll_rows()
    if not rows:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll Report"

    # ── Header row ──
    headers = ["Emp ID", "Name", "Department", "Designation",
               "Basic Salary", "Bonus", "OT Pay", "Gross Salary",
               "Deductions", "Net Salary", "Attendance", "OT Hrs"]

    header_fill = PatternFill("solid", fgColor="4F46E5")   # indigo
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # ── Data rows ──
    for row_idx, p in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1,  value=p["emp_id"])
        ws.cell(row=row_idx, column=2,  value=p["name"])
        ws.cell(row=row_idx, column=3,  value=p["department"])
        ws.cell(row=row_idx, column=4,  value=p["designation"])
        ws.cell(row=row_idx, column=5,  value=p["basic_salary"])
        ws.cell(row=row_idx, column=6,  value=p["bonus"])
        ws.cell(row=row_idx, column=7,  value=p["overtime_pay"])
        ws.cell(row=row_idx, column=8,  value=p["gross_salary"])
        ws.cell(row=row_idx, column=9,  value=p["deductions"])
        ws.cell(row=row_idx, column=10, value=p["net_salary"])
        ws.cell(row=row_idx, column=11, value=p["attendance"])
        ws.cell(row=row_idx, column=12, value=p["overtime_hrs"])

        # Alternate row shading
        if row_idx % 2 == 0:
            light_fill = PatternFill("solid", fgColor="EEF2FF")
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = light_fill

    # ── Auto-fit columns ──
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    # ── Summary sheet ──
    ws2 = wb.create_sheet("Summary")
    total_payroll = sum(p["net_salary"] for p in rows)
    avg_salary    = total_payroll / len(rows)
    highest       = max(rows, key=lambda x: x["net_salary"])

    summary_data = [
        ("Total Employees",      len(rows)),
        ("Total Payroll Expense",f"₹ {total_payroll:,.2f}"),
        ("Average Net Salary",   f"₹ {avg_salary:,.2f}"),
        ("Highest Paid Employee",f"{highest['name']} (₹ {highest['net_salary']:,.2f})"),
        ("Export Date",          datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for r, (label, value) in enumerate(summary_data, start=1):
        ws2.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws2.cell(row=r, column=2, value=value)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename  = f"payroll_report_{timestamp}.xlsx"
    filepath  = os.path.join(EXPORTS_DIR, filename)
    wb.save(filepath)
    print(f"\n✅  Excel exported → {filepath}")
